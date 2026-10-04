import io

from openpyxl import load_workbook

from tools.make_dummy_pdfs import Spec, build_pdf

from .conftest import pdf_files


def test_import_endpoint_and_lists(client):
    res = client.post("/api/import", files=pdf_files()).json()["results"]
    assert [r["status"] for r in res] == ["ok"] * 3
    ts = client.get("/api/tariffs").json()
    assert {t["origin"]: t["status"] for t in ts} == {"拠点A": "active", "拠点B": "active", "拠点C": "expired"}
    assert client.get("/api/tariffs/999").status_code == 404


def test_bad_upload_reports_kind_only(client):
    r = client.post("/api/import", files=[("files", ("secret_customer.pdf", b"not pdf", "application/pdf"))]).json()["results"][0]
    assert r == {"filename": "secret_customer.pdf", "status": "error", "error": "not_a_pdf"}


def test_filename_path_is_stripped(client):
    r = client.post("/api/import", files=[("files", ("../../etc/x.pdf", build_pdf(Spec("拠点Z")), "application/pdf"))]).json()["results"][0]
    assert r["filename"] == "x.pdf"


def test_lookup_combined_notes(loaded):
    d = loaded.get("/api/lookup", params={"dest": "大阪", "size": 80}).json()
    assert d["destination"]["region"] == "関西" and len(d["results"]) == 3
    c = loaded.get("/api/combined").json()
    assert len(c["rows"]) == 3 * 132 and c["cheapest"]
    n = loaded.get("/api/notes").json()
    assert any("沖縄" in x["text"] for x in n[0]["notes"])


def test_manual_edit_api(loaded):
    tid = [t for t in loaded.get("/api/tariffs").json() if t["origin"] == "拠点C"][0]["id"]
    r = loaded.put(f"/api/tariffs/{tid}/rates", json={"region": "関東", "size": 100, "price": 555})
    assert r.json()["rates"]["関東"]["100"]["price"] == 555
    assert loaded.put(f"/api/tariffs/{tid}/rates", json={"region": "関東", "size": 100, "price": 0}).status_code == 422
    assert loaded.put(f"/api/tariffs/{tid}", json={"origin": "拠点C改"}).json()["origin"] == "拠点C改"
    assert loaded.delete(f"/api/tariffs/{tid}").status_code == 200


def test_exports(loaded):
    csv = loaded.get("/api/export.csv")
    assert csv.headers["content-type"].startswith("text/csv") and csv.content.startswith(b"\xef\xbb\xbf")
    assert len(csv.text.strip().splitlines()) == 1 + 3 * 132
    wb = load_workbook(io.BytesIO(loaded.get("/api/export.xlsx").content))
    assert wb.sheetnames[:2] == ["統合", "比較"] and {"拠点A", "拠点B", "拠点C", "注記"} <= set(wb.sheetnames)
    assert wb["統合"].max_row == 1 + 3 * 132
    assert wb["拠点A"].max_row == 3 + 12


def test_static_and_no_openapi_cdn(client):
    assert "運賃表ハブ" in client.get("/").text
    assert client.get("/static/app.js").status_code == 200
    for p in ("/docs", "/redoc", "/openapi.json"):  # Swagger UI は外部CDNを読むため無効
        assert client.get(p).status_code == 404
    assert "default-src 'self'" in client.get("/").headers["content-security-policy"]


def test_host_and_origin_checks(client):
    assert client.get("/api/tariffs", headers={"host": "evil.example"}).status_code == 400
    r = client.post("/api/import", files=pdf_files()[:1], headers={"origin": "http://evil.example"})
    assert r.status_code == 403
    r = client.post("/api/import", files=pdf_files()[:1], headers={"origin": "http://testserver"})
    assert r.status_code == 200


def test_optional_password(tmp_path):
    from fastapi.testclient import TestClient
    from app.main import create_app
    c = TestClient(create_app(db_path=tmp_path / "p.sqlite", password="pw", allowed_hosts={"testserver"}))
    assert c.get("/api/tariffs").status_code == 401
    assert c.get("/api/tariffs", auth=("u", "wrong")).status_code == 401
    assert c.get("/api/tariffs", auth=("u", "pw")).status_code == 200
