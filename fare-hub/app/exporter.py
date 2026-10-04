"""CSV / Excel 出力。ファイルは書かずメモリ上で生成して返す（exports/ もGit管理外）。"""
from __future__ import annotations

import csv
import io
import re
from datetime import date

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from . import service
from .config import Settings

STATUS_LABEL = {"active": "有効", "expiring": "期限間近", "expired": "期限切れ", "upcoming": "適用前", "unknown": "期日不明"}
HEADERS = ["発地拠点", "着地地帯", "サイズ", "重量目安(kg)", "運賃(税別・円)", "区分", "運賃表状態",
           "適用期日", "満期日", "見積No.", "顧客コード", "最安"]
GREEN = PatternFill("solid", fgColor="C6EFCE")
GREY = PatternFill("solid", fgColor="D9D9D9")


def _safe(v):
    """表計算ソフトの数式として解釈されうる文字列を無害化（CSV/Excelインジェクション対策）。"""
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + v
    return v


def _flat(rows: list[dict], best: dict) -> list[list]:
    out = []
    for r in rows:
        is_best = r["price"] is not None and best.get((r["region"], r["size"])) == r["price"] and r["status"] not in ("expired", "upcoming")
        out.append([r["origin"], r["region"], r["size"], r["weight_kg"], r["price"],
                    {"parsed": "PDF取込", "manual": "手動補正"}.get(r["source"], "欠損"),
                    STATUS_LABEL[r["status"]], r["valid_from"], r["valid_to"], r["quote_no"], r["customer_code"],
                    "★" if is_best else ""])
    return out


def to_csv(conn, st: Settings, today: date) -> bytes:
    rows = service.combined_rows(conn, st, today)
    buf = io.StringIO()
    w = csv.writer(buf)
    w.writerow(HEADERS)
    for line in _flat(rows, service.cheapest(rows)):
        w.writerow([_safe(v) if v is not None else "" for v in line])
    return ("﻿" + buf.getvalue()).encode("utf-8")  # BOM付き（Excelで文字化けしない）


def _sheet_name(wb: Workbook, base: str) -> str:
    name = re.sub(r"[\[\]:*?/\\]", "_", base)[:28] or "sheet"
    cand, i = name, 2
    while cand in wb.sheetnames:
        cand = f"{name[:25]}_{i}"
        i += 1
    return cand


def to_xlsx(conn, st: Settings, today: date) -> bytes:
    rows = service.combined_rows(conn, st, today)
    best = service.cheapest(rows)
    wb = Workbook()

    ws = wb.active
    ws.title = "統合"
    ws.append(HEADERS)
    for line in _flat(rows, best):
        ws.append([_safe(v) for v in line])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"

    # 比較: 地帯 × サイズ の行、拠点ごとの列。最安セルを緑に。
    origins = sorted({r["origin"] for r in rows})
    cmp_ws = wb.create_sheet("比較")
    cmp_ws.append(["着地地帯", "サイズ", "最安運賃", "最安拠点"] + [_safe(o) for o in origins])
    price = {(r["origin"], r["region"], r["size"]): r for r in rows}
    for region in st.region_names:
        for size in st.sizes:
            b = best.get((region, size))
            cells = [price.get((o, region, size)) for o in origins]
            winners = [o for o, c in zip(origins, cells)
                       if c and c["price"] == b and c["status"] not in ("expired", "upcoming")] if b is not None else []
            cmp_ws.append([region, size, b, "・".join(winners) if winners else None] + [c["price"] if c else None for c in cells])
            for j, (o, c) in enumerate(zip(origins, cells)):
                if c and c["price"] is not None:
                    cell = cmp_ws.cell(row=cmp_ws.max_row, column=5 + j)
                    if o in winners:
                        cell.fill = GREEN
                    elif c["status"] in ("expired", "upcoming"):
                        cell.fill = GREY
    cmp_ws.freeze_panes = "C2"

    # 拠点別マトリクス
    for t in service.current_tariffs(conn, st, today):
        d = service.get_tariff(conn, st, t["id"], today)
        s = wb.create_sheet(_sheet_name(wb, t["origin"]))
        s.append([_safe(t["origin"]), f"{STATUS_LABEL[t['status']]}  {t['valid_from'] or '?'} 〜 {t['valid_to'] or '?'}"])
        s.append(["着地＼サイズ"] + st.sizes)
        s.append(["重量目安(kg)"] + [d["weights"].get(str(z)) for z in st.sizes])
        for region in st.region_names:
            s.append([region] + [(d["rates"][region].get(str(z)) or {}).get("price") for z in st.sizes])
        for c in s[2] + s[3]:
            c.font = Font(bold=True)

    # 注記
    ns = wb.create_sheet("注記")
    ns.append(["発地拠点", "分類", "注記"])
    labels = {c["key"]: c["label"] for c in st.parser["notes"]["categories"]}
    for t in service.current_tariffs(conn, st, today):
        for n in service.get_tariff(conn, st, t["id"], today)["notes"]:
            ns.append([_safe(t["origin"]), labels.get(n["category"], "その他"), _safe(n["text"])])

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
