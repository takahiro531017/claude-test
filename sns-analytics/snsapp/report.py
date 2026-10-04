"""分析ファイル(HTML / PDF / Excel)の作成。"""
import re
import shutil
from datetime import datetime
from pathlib import Path

from jinja2 import Environment, FileSystemLoader, select_autoescape
from openpyxl import Workbook
from openpyxl.chart import LineChart, Reference
from openpyxl.styles import Alignment, Font, PatternFill

from . import db

HERE = Path(__file__).resolve().parent
CSS = (HERE / "style.css").read_text(encoding="utf-8")
_env = Environment(loader=FileSystemLoader(HERE.parent / "templates"), autoescape=select_autoescape(["html"]))


def _macros():
    return _env.get_template("report.html").module


def render_section(name, a, *args) -> str:
    """画面表示用。<style>つきの部品HTMLを返す。"""
    return f"<style>{CSS}</style><div class='card-wrap'>" + str(getattr(_macros(), name)(a, *args)) + "</div>"


def render_html(a) -> str:
    return str(_macros().full(a, CSS))


def render_pdf(html_path: Path, pdf_path: Path):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        try:
            b = p.chromium.launch()
        except Exception:
            exe = shutil.which("chromium") or "/opt/pw-browsers/chromium"
            b = p.chromium.launch(executable_path=exe)
        page = b.new_page()
        page.goto(html_path.resolve().as_uri())
        page.pdf(path=str(pdf_path), format="A4", print_background=True, margin=dict(top="12mm", bottom="12mm", left="10mm", right="10mm"))
        b.close()


HEAD = Font(bold=True, color="FFFFFF")
FILL = PatternFill("solid", fgColor="2B6CB0")


def _sheet(wb, title, header, rows, widths=None):
    ws = wb.create_sheet(title)
    ws.append(header)
    for c in ws[1]:
        c.font, c.fill = HEAD, FILL
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    return ws


def render_xlsx(a, path: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = "今回のまとめ"
    ws.append([f"{a['client']} / Instagram / 対象期間 {a['period']} / 作成 {a['created']}"])
    ws["A1"].font = Font(bold=True, size=13)
    ws.append([])
    for i, s in enumerate(a["summary"], 1):
        ws.append([f"{i}. {s}"])
    ws.column_dimensions["A"].width = 120
    _sheet(wb, "数字の変化", ["項目", "意味", "今回"] + [f"{n}と比べて" for n in a["ref_names"]],
           [[r["label"], "→ " + r["short"], r["current"]] + [f"{c['text']}(相手: {c['ref']})" for c in r["cells"]] for r in a["changes"]],
           [22, 34, 14, 28, 28, 28])
    rows = []
    if a["posts"]["ok"]:
        for tag, key in [("ベスト", "top"), ("ワースト", "worst")]:
            for i, p in enumerate(a["posts"][key], 1):
                rows.append([f"{tag}{i}", p["date"], p["type"], p["theme"], p["rate"], p["reach"], p["likes"], p["saves"], p["why"]])
    else:
        rows.append([a["posts"]["note"]])
    _sheet(wb, "投稿ベスト・ワースト", ["順位", "投稿日時", "種類", "テーマ", "エンゲージメント率→反応した人の割合", "リーチ→見た人の数", "いいね", "保存", "理由の仮説"],
           rows, [10, 18, 12, 14, 18, 14, 10, 10, 80])
    for b in a["breakdowns"]:
        _sheet(wb, b["title"], ["区分", "投稿数", "平均リーチ→見た人の数", "エンゲージメント率→反応した人の割合", "保存率→保存した人の割合"],
               [[r["key"], r["n"], None if r["avg_reach"] is None else round(r["avg_reach"]), r["rate"], r["save_rate"]] for r in b.get("rows", [])]
               or [["データなし"]], [20, 10, 22, 24, 22])
        for row in wb[b["title"]].iter_rows(min_row=2, min_col=4, max_col=5):
            for c in row:
                c.number_format = "0.00%"
    s = a["series"]
    fw = _sheet(wb, "フォロワー推移", ["期間", "フォロワー数→フォローしている人の数", "フォロワー増加数→期間中に増えた人数"],
                [[l, f, g] for l, f, g in zip(s["labels"], s["followers"], s["gain"])], [14, 30, 30])
    if len(s["labels"]) > 1:
        lc = LineChart()
        lc.title, lc.height, lc.width = "フォロワー数の推移", 8, 18
        lc.add_data(Reference(fw, min_col=2, min_row=1, max_row=len(s["labels"]) + 1), titles_from_data=True)
        lc.set_categories(Reference(fw, min_col=1, min_row=2, max_row=len(s["labels"]) + 1))
        fw.add_chart(lc, "E2")
    _sheet(wb, "深掘り分析", ["テーマ", "種類", "内容"],
           [[d["title"], "数字の言葉", f"{n} → {a['glossary_map'][n]}"] for d in a["deep"] for n in d["terms"]]
           + [[d["title"], "グラフのコメント", c["comment"]] for d in a["deep"] for c in d["charts"]]
           + [[d["title"], "理由の仮説", h] for d in a["deep"] for h in d["hypotheses"]], [30, 16, 110])
    _sheet(wb, "次にやること", ["#", "やること", "なぜ", "どうやる"],
           [[i, x["title"], x["why"], x["how"]] for i, x in enumerate(a["suggestions"], 1)], [5, 36, 70, 70])
    _sheet(wb, "用語集", ["用語", "→ 意味", "→ 数字の見方(良い目安)"], [[g["term"], "→ " + g["meaning"], g["guide"]] for g in a["glossary"]], [22, 60, 80])
    wb.save(path)


def _slug(s):
    return re.sub(r"[^\w\-]+", "_", s)[:40].strip("_") or "client"


def create_files(client_id, a, label="", make_pdf=True):
    """HTML・Excel・PDFを output/ に作り、履歴に記録する。戻り値: dict(paths..., pdf_error)"""
    folder = db.output_dir() / f"client_{client_id}_{_slug(a['client'])}"
    folder.mkdir(parents=True, exist_ok=True)
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    base = folder / f"{stamp}_{_slug(a['period_label'])}{('_' + _slug(label)) if label else ''}"
    html, xlsx, pdf = base.with_suffix(".html"), base.with_suffix(".xlsx"), base.with_suffix(".pdf")
    html.write_text(render_html(a), encoding="utf-8")
    render_xlsx(a, xlsx)
    err = None
    if make_pdf:
        try:
            render_pdf(html, pdf)
        except Exception as e:  # noqa: BLE001
            err = str(e).splitlines()[0] if str(e) else "PDFを作れませんでした"
            pdf = None
    else:
        pdf = None
    db.add_report(client_id, a["created"], a["period"], label, str(html), str(pdf) if pdf else "", str(xlsx))
    return dict(html=html, xlsx=xlsx, pdf=pdf, pdf_error=err)
