from datetime import date

import openpyxl

from shiire_check import excel_io
from shiire_check.cache import Cache
from shiire_check.matcher import ResultRow
from shiire_check.schema import ReadResult


def test_ledger_roundtrip(dummy):
    path = dummy["ledger"]
    assert excel_io.list_sheets(path) == ["仕入一覧"]
    heads = excel_io.read_headers(path, "仕入一覧", 1)
    assert heads[2] == "C列: 伝票番号"
    rows = excel_io.read_ledger(path, "仕入一覧", 1, number_col=2, maker_col=1, date_col=0, amount_col=3)
    assert len(rows) == 13
    assert rows[3].number == "4567"           # 数値セルが 4567.0 にならない
    assert rows[0].row == 2 and rows[0].maker == "山田製作所"


def test_export(tmp_path):
    rows = [
        ResultRow("k1", "一致", "100", "A社", "2026-10-01", "備考", "a.jpg", "高", warnings=["金額相違"]),
        ResultRow("k2", "要確認", "=cmd()", "", "", "手書き", "b.jpg", "低", confirmed=True, memo="確認した"),
        ResultRow("ledger:1#1", "伝票なし", "300", "", "", "", "", "", is_slip=False),
    ]
    name = excel_io.default_output_name(date(2026, 10, 5))
    assert name == "仕入チェック結果_2026-10-05.xlsx"
    out = excel_io.write_results(tmp_path / name, rows, today=date(2026, 10, 5))
    wb = openpyxl.load_workbook(out)
    assert wb.sheetnames == ["照合結果", "集計"]
    ws = wb["照合結果"]
    assert [c.value for c in ws[1]][:8] == ["日付", "区分", "伝票番号", "メーカー名", "元ファイル名", "自信度", "確認済み", "担当者メモ"]
    assert ws["B2"].value == "一致（金額相違）"
    assert ws["C3"].value == "=cmd()" and ws["C3"].data_type == "s"   # 数式にならない
    assert ws["G3"].value == "済" and ws["H3"].value == "確認した"
    assert ws["A4"].value == "2026-10-05"
    agg = {r[0].value: r[1].value for r in wb["集計"].iter_rows(min_row=1)}
    assert agg["一致"] == 1 and agg["要確認"] == 1 and agg["伝票なし"] == 1 and agg["一覧にない"] == 0
    assert agg["合計"] == 3 and agg["確認済み"] == 1 and agg["警告あり（金額相違・重複など）"] == 1


def test_cache(tmp_path):
    c = Cache(tmp_path / "c.db")
    r = ReadResult(slip_numbers=["1"], handwritten=False, confidence="high")
    assert c.get_reading("k", "m") is None
    c.put_reading("k", "m", r)
    assert c.get_reading("k", "m") == r
    assert c.get_reading("k", "other-model") is None
    c.set_correction("k", "99")
    assert c.get_corrections() == {"k": "99"}
    c.set_correction("k", "")
    assert c.get_corrections() == {}
    c.set_state("k", True, "メモ")
    c.close()
    assert Cache(tmp_path / "c.db").get_states() == {"k": (True, "メモ")}   # 再起動しても残る
