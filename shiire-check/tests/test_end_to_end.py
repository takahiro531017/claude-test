"""ダミーデータ一式（モック読み取り）で、4区分が期待どおりに分類されることを確認する。"""
import threading

import pytest

from shiire_check.cache import Cache
from shiire_check.matcher import MATCH, NO_SLIP, NOT_IN_LEDGER, REVIEW
from shiire_check.reader import MockReader
from shiire_check.service import CheckSession, ColumnSelection

EXPECT_SLIPS = {
    "山田製作所_001.jpg": (MATCH, ["伝票番号重複"]),
    "山田製作所_001_再スキャン.jpg": (MATCH, ["伝票番号重複"]),
    "東海部品_a.png": (MATCH, []),           # TK-00123 == TK-123
    "東海部品_b.png": (MATCH, []),           # 候補2つのうち一覧にあるものを採用
    "東海部品_c.png": (REVIEW, []),          # 候補が複数で特定不能
    "北斗化学.jpg": (MATCH, []),             # 全角・先頭ゼロ・数値セル
    "丸山商会_手書き.png": (REVIEW, []),     # 手書き
    "大和電機_まとめ.pdf（1ページ目）": (MATCH, []),
    "大和電機_まとめ.pdf（2ページ目）": (NOT_IN_LEDGER, []),
    "大和電機_まとめ.pdf（3ページ目）": (MATCH, ["金額相違"]),
    "大阪精機_横向き.jpg": (MATCH, []),
    "三光産業_ぼやけ.jpg": (REVIEW, []),     # 読み取れない
    "青葉工業_かすれ.png": (REVIEW, []),     # 自信度が低い
    "松本商事_失敗.png": (REVIEW, []),       # API障害
    "小林工業.png": (NOT_IN_LEDGER, []),
    "壊れたファイル.jpg": (REVIEW, []),
}


@pytest.fixture
def session(cfg, dummy, tmp_path):
    s = CheckSession(cfg, Cache(tmp_path / "c.db"))
    s.load_ledger(dummy["ledger"], ColumnSelection("仕入一覧", 1, number=2, maker=1, date=0, amount=3))
    s.read_slips(dummy["slips"], MockReader(dummy["slips"]))
    return s


def test_categories(session):
    got = {r.source_file: r for r in session.results if r.is_slip}
    assert set(got) == set(EXPECT_SLIPS)
    for name, (cat, warns) in EXPECT_SLIPS.items():
        assert (got[name].category, got[name].warnings) == (cat, warns), (name, got[name].note)
    no_slip = sorted(r.slip_number for r in session.results if r.category == NO_SLIP)
    # 本当に伝票が無い2件 + 番号が全く読めなかった伝票（ぼやけ・API障害・壊れたファイル）の分は「伝票なし」に残る
    assert no_slip == ["MS-4410", "NM-0001", "NM-0002", "SK-3319"]
    assert all("読み取れなかった伝票が3枚" in r.note for r in session.results if r.category == NO_SLIP)
    # 番号が読めている要確認（手書き・低自信度・候補複数）の一覧行は「伝票なし」にならない
    assert "H2210" not in no_slip and "AB-5521" not in no_slip


def test_read_failure_reasons(session):
    notes = {r.source_file: r.note for r in session.results if r.is_slip}
    assert "読み取り失敗" in notes["松本商事_失敗.png"]
    assert "自信度が低" in notes["青葉工業_かすれ.png"]
    assert "手書き" in notes["丸山商会_手書き.png"]


def test_correction_rematches_and_persists(session, cfg, dummy, tmp_path):
    row = next(r for r in session.results if r.source_file == "三光産業_ぼやけ.jpg")
    assert row.category == REVIEW
    session.set_correction(row.key, "SK-3319")             # 人が目で見て入力
    row = next(r for r in session.results if r.source_file == "三光産業_ぼやけ.jpg")
    assert row.category == MATCH and "手修正済み" in row.note
    assert not any(r.slip_number == "SK-3319" and r.category == NO_SLIP for r in session.results)   # 修正後は伝票なしも消える
    # 確認済みもSQLiteに残り、別セッション（アプリ再起動）でも復元される
    session.set_state(row, True, "現物確認")
    s2 = CheckSession(cfg, Cache(tmp_path / "c.db"))
    s2.ledger, s2.cols, s2.readings = session.ledger, session.cols, session.readings
    again = next(r for r in s2.rematch() if r.source_file == "三光産業_ぼやけ.jpg")
    assert again.category == MATCH and again.confirmed and again.memo == "現物確認"
    # 修正の取り消し
    s2.set_correction(row.key, "")
    assert next(r for r in s2.results if r.source_file == "三光産業_ぼやけ.jpg").category == REVIEW


def test_export(session, tmp_path):
    import openpyxl
    out = session.export(tmp_path / "out.xlsx")
    ws = openpyxl.load_workbook(out)["照合結果"]
    labels = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
    assert ws.max_row - 1 == len(session.results)
    assert "一致（金額相違）" in labels and "要確認" in labels and "伝票なし" in labels and "一覧にない" in labels
