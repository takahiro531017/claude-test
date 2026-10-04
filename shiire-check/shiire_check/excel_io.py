"""仕入一覧表の読み込みと、照合結果のExcel出力。"""
from __future__ import annotations

from datetime import date, datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from shiire_check.matcher import CATEGORIES, LedgerRow, ResultRow, summarize


def _text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    if isinstance(v, (datetime, date)):
        return v.isoformat()
    return str(v).strip()


def list_sheets(path: str | Path) -> list[str]:
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        return list(wb.sheetnames)
    finally:
        wb.close()


def read_headers(path: str | Path, sheet: str | None = None, header_row: int = 1) -> list[str]:
    """ヘッダー行のセルを「A列: 伝票番号」の形で返す（列名は固定しない）。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]
        row = next(ws.iter_rows(min_row=header_row, max_row=header_row, values_only=True), ())
    finally:
        wb.close()
    heads = []
    for i, v in enumerate(row, start=1):
        name = _text(v)
        heads.append(f"{get_column_letter(i)}列: {name}" if name else f"{get_column_letter(i)}列: （見出しなし）")
    return heads


def read_ledger(
    path: str | Path,
    sheet: str | None,
    header_row: int,
    number_col: int,
    maker_col: int | None = None,
    date_col: int | None = None,
    amount_col: int | None = None,
) -> list[LedgerRow]:
    """列は0始まりの位置で指定。伝票番号が空の行は除く。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    rows: list[LedgerRow] = []
    try:
        ws = wb[sheet] if sheet else wb.worksheets[0]

        def cell(vals, idx):
            return vals[idx] if idx is not None and idx < len(vals) else None

        for r, vals in enumerate(ws.iter_rows(min_row=header_row + 1, values_only=True), start=header_row + 1):
            num = _text(cell(vals, number_col))
            if not num:
                continue
            rows.append(LedgerRow(
                row=r, number=num,
                maker=_text(cell(vals, maker_col)) or None,
                date=cell(vals, date_col), amount=cell(vals, amount_col),
            ))
    finally:
        wb.close()
    return rows


FILL = {
    "一致": "D9F0DD", "伝票なし": "FFF7B3", "一覧にない": "F8C9C9", "要確認": "FFD9A8",
}
HEADERS = ["日付", "区分", "伝票番号", "メーカー名", "元ファイル名", "自信度", "確認済み", "担当者メモ", "備考"]


def default_output_name(today: date | None = None) -> str:
    return f"仕入チェック結果_{(today or date.today()).isoformat()}.xlsx"


def _put(ws, r, c, v):
    cell = ws.cell(row=r, column=c, value=v)
    if isinstance(v, str) and v[:1] in "=+-@":   # 数式として解釈されないよう文字列固定
        cell.data_type = "s"
    return cell


def write_results(path: str | Path, rows: list[ResultRow], today: date | None = None) -> Path:
    today = today or date.today()
    wb = Workbook()
    ws = wb.active
    ws.title = "照合結果"
    for c, h in enumerate(HEADERS, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="305496")
    for r, row in enumerate(rows, start=2):
        vals = [
            row.date or today.isoformat(), row.label, row.slip_number, row.maker, row.source_file,
            row.confidence, "済" if row.confirmed else "", row.memo, row.note,
        ]
        for c, v in enumerate(vals, start=1):
            cell = _put(ws, r, c, v)
            cell.fill = PatternFill("solid", fgColor=FILL[row.category])
            cell.alignment = Alignment(vertical="top", wrap_text=(c == 9))
    for c, w in enumerate([12, 20, 22, 22, 28, 8, 10, 24, 60], start=1):
        ws.column_dimensions[get_column_letter(c)].width = w
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{max(1, len(rows) + 1)}"

    ws2 = wb.create_sheet("集計")
    ws2.append(["処理日", today.isoformat()])
    ws2.append([])
    ws2.append(["区分", "件数"])
    counts = summarize(rows)
    for cat in CATEGORIES:
        ws2.append([cat, counts[cat]])
    ws2.append(["合計", len(rows)])
    ws2.append([])
    ws2.append(["警告あり（金額相違・重複など）", sum(1 for r in rows if r.warnings)])
    ws2.append(["確認済み", sum(1 for r in rows if r.confirmed)])
    ws2.append(["未確認の「要確認」", sum(1 for r in rows if r.category == "要確認" and not r.confirmed)])
    for cell in ws2[3]:
        cell.font = Font(bold=True)
    ws2.column_dimensions["A"].width = 34

    path = Path(path)
    wb.save(path)
    return path
